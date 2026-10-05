#!/usr/bin/env python3
"""
Objective checks for the eval cases.

Usage:
    check_outputs.py <case_dir> <outputs_dir> [--json]

<outputs_dir> holds the files one run produced (PDF, .txt, .docx, ...). Checks
come from <case_dir>/case.json. The CV is the PDF (or .txt/.docx) whose name does
not look like a cover letter; the letter is the one whose name matches
cover|personligt|brev|letter. Needs pdftotext/pdfinfo (poppler); python-docx for
.docx. Exit code 0 if every check passes.
"""

import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

LETTER_RE = re.compile(r"cover|personligt|brev|letter", re.I)
EM_DASH = "—"


def pdf_text(path):
    exe = shutil.which("pdftotext")
    if not exe:
        raise SystemExit("pdftotext (poppler) is required for the eval checks")
    return subprocess.run([exe, "-layout", str(path), "-"], capture_output=True, text=True).stdout


def pdf_pages(path):
    exe = shutil.which("pdfinfo")
    out = subprocess.run([exe, str(path)], capture_output=True, text=True).stdout if exe else ""
    m = re.search(r"Pages:\s+(\d+)", out)
    return int(m.group(1)) if m else None


def docx_text(path):
    from docx import Document
    d = Document(str(path))
    return "\n".join(p.text for p in d.paragraphs), d


def collapse_spaced(text):
    """Letter-spaced headings ('E R FA R E N H E T') come out of PDFs with gaps; rejoin them."""
    out = []
    for line in text.splitlines():
        t = line.strip()
        if re.fullmatch(r"(?:\S ){3,}\S|(?:\S{1,2} ){3,}\S{1,2}", t) and len(t.replace(" ", "")) <= 24:
            t = t.replace(" ", "")
        out.append(t)
    return "\n".join(out)


NEGATION_RE = re.compile(r"\b(inte|aldrig|saknar|utan|ingen|inget|not|never|no|without|lack|haven't|don't|hasn't)\b", re.I)
APPLICATION_RE = re.compile(r"(ansökan|söker|sökande|application|applying)\s+(till|to|for|som)?\s*$", re.I)


def find(outputs, suffix, want_letter):
    hits = [p for p in sorted(outputs.rglob(f"*{suffix}")) if bool(LETTER_RE.search(p.name)) == want_letter]
    return hits[0] if hits else None


def run(case_dir, outputs):
    case = json.loads((case_dir / "case.json").read_text())
    results = []

    def check(name, ok, detail=""):
        results.append({"check": name, "pass": bool(ok), "detail": detail})

    cv_pdf = find(outputs, ".pdf", False)
    cv_txt = find(outputs, ".txt", False)
    cv_docx = find(outputs, ".docx", False)
    letter_pdf, letter_txt = find(outputs, ".pdf", True), find(outputs, ".txt", True)
    letter_docx = find(outputs, ".docx", True)

    check("PDF CV delivered", cv_pdf is not None, str(cv_pdf.name) if cv_pdf else "no CV PDF found")
    cv_text = collapse_spaced(pdf_text(cv_pdf)) if cv_pdf else (cv_txt.read_text() if cv_txt else "")
    if cv_pdf:
        n = pdf_pages(cv_pdf)
        lo, hi = case["pages"]
        check(f"{lo}-{hi} pages", n is not None and lo <= n <= hi, f"{n} pages")

    texts = {"CV": cv_text}
    if letter_pdf:
        texts["letter"] = pdf_text(letter_pdf)
    elif letter_txt:
        texts["letter"] = letter_txt.read_text()
    for p in (cv_txt, letter_txt):
        if p:
            texts[p.name] = p.read_text()
    if cv_docx:
        try:
            texts[cv_docx.name] = docx_text(cv_docx)[0]
        except Exception as e:
            check("DOCX opens", False, str(e))

    check("no em dashes", all(EM_DASH not in t for t in texts.values()),
          ", ".join(k for k, t in texts.items() if EM_DASH in t))

    leaks = [(k, w) for k, t in texts.items() for w in case["forbidden_personal"] if w.lower() in t.lower()]
    check("no DOB / marital status / address / personnummer / photo note", not leaks, str(leaks))

    invented = []
    for skill in case["forbidden_skills"]:
        if re.search(re.escape(skill), cv_text, re.I):
            invented.append(skill)
        letter = texts.get("letter", "")
        claim = re.compile(r"(erfarenhet av|arbetat (med|i)|använt|använder|kunskaper i|experience (with|of|in)|worked (with|in)|skilled in)\s+" + re.escape(skill), re.I)
        for m in claim.finditer(letter):
            before = letter[max(0, m.start() - 60):m.start()].lower()
            # "Jag har inte arbetat med X" / "no experience of X" are honest gap statements, not claims.
            if not NEGATION_RE.search(before):
                invented.append(f"{skill} (claimed in letter)")
    check("no invented skills " + "/".join(case["forbidden_skills"]), not invented, str(invented))

    header = "\n".join([l for l in cv_text.splitlines() if l.strip()][:4])
    bad = []
    for pat in case["header_title_must_not_match"]:
        for m in re.finditer(pat, header, re.I):
            # "Application to Customer Success Manager" names the target role; it does not claim the title.
            if not APPLICATION_RE.search(header[max(0, m.start() - 20):m.start()]):
                bad.append(pat)
    check("honest header title", not bad, f"header: {header!r}" if bad else "")

    missing = [k for k in case["must_keywords"] if k.lower() not in cv_text.lower()]
    check("must-have keywords present", not missing, f"missing: {missing}" if missing else "")

    if case["swedish_headings"]:
        lines = {l.strip().lower() for l in cv_text.splitlines()}
        has = lambda pat: any(re.fullmatch(pat, l) for l in lines)
        check("Swedish section headings",
              has(r"\w*erfarenhet\w*") and has(r"utbildning") and not ({"experience", "education", "work experience"} & lines),
              "headings seen: " + ", ".join(sorted(l for l in lines if 3 < len(l) < 24 and l.isalpha()))[:200])
    else:
        lines = {l.strip().lower() for l in cv_text.splitlines()}
        check("English section headings", "experience" in lines or "work experience" in lines)

    if case["expect_docx"]:
        ok = False
        detail = "no .docx found"
        if cv_docx:
            try:
                text, d = docx_text(cv_docx)
                ok = len(text.split()) > 150 and len(d.tables) == 0
                detail = f"{len(text.split())} words, {len(d.tables)} tables"
            except Exception as e:
                detail = str(e)
        check(".docx delivered for portal/ATS", ok, detail)

    if case["expect_cover_letter"]:
        lo, hi = case["cover_letter_words"]
        t = texts.get("letter", "")
        body = re.sub(r"\s+", " ", t)
        n = len(body.split())
        check(f"cover letter delivered, {lo}-{hi} words", bool(t) and lo <= n <= hi, f"{n} words")
    return results


def main():
    args = [a for a in sys.argv[1:] if a != "--json"]
    if len(args) != 2:
        print(__doc__)
        sys.exit(2)
    results = run(Path(args[0]), Path(args[1]))
    if "--json" in sys.argv:
        print(json.dumps(results, indent=2, ensure_ascii=False))
    else:
        for r in results:
            print(("PASS " if r["pass"] else "FAIL ") + r["check"] + (f"  [{r['detail']}]" if r["detail"] and not r["pass"] else ""))
        print(f"{sum(r['pass'] for r in results)}/{len(results)} checks passed")
    sys.exit(0 if all(r["pass"] for r in results) else 1)


if __name__ == "__main__":
    main()
