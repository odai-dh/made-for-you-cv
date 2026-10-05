#!/usr/bin/env python3
"""
Smoke test: fill every template with fictional sample data, build PDFs, and
check the output. Run from the repo root:  python3 tests/smoke_test.py
Exits non-zero on failure. Needs one PDF renderer (Chrome/Chromium, WeasyPrint
or Playwright).
"""
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "skills" / "made-for-you-cv" / "scripts"
TESTS = Path(__file__).resolve().parent
failures = []


def run(*args):
    return subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True)


def check(cond, msg):
    print(("  ok   " if cond else "  FAIL ") + msg)
    if not cond:
        failures.append(msg)


with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    cv = json.loads((TESTS / "sample_cv.json").read_text())

    jobs = [(t, TESTS / "sample_cv.json", 2) for t in ("minimal", "two-column", "designer-accent")]
    jobs.append(("cover-letter", TESTS / "sample_cover_letter.json", 1))

    for tpl, data, max_pages in jobs:
        print(f"\n[{tpl}]")
        html_out, pdf_out, txt_out = tmp / f"{tpl}.html", tmp / f"{tpl}.pdf", tmp / f"{tpl}.txt"
        r = run(SCRIPTS / "fill_template.py", tpl, data, html_out)
        check(r.returncode == 0, "fill_template succeeds " + r.stderr.strip())
        if r.returncode:
            continue
        html = html_out.read_text()
        check("{{" not in html and "<!--IF" not in html, "no leftover tokens or markers")
        r = run(SCRIPTS / "build_cv.py", html_out, pdf_out)
        check(r.returncode == 0, "build_cv produces a PDF " + r.stderr.strip()[-200:])
        if r.returncode == 0:
            m = re.search(r"Page count: (\d+)", r.stdout)
            check(bool(m) and int(m.group(1)) <= max_pages, f"page count {m.group(1) if m else '?'} <= {max_pages}")
        r = run(SCRIPTS / "html_to_text.py", html_out, txt_out)
        txt = txt_out.read_text() if txt_out.exists() else ""
        check("Jane Doe" in txt, "plain-text export contains the name")

    print("\n[optional fields]")
    slim = {k: v for k, v in cv.items() if k not in ("PHONE", "PORTFOLIO", "GITHUB", "PROJECTS_BLOCK", "LANGUAGES_BLOCK")}
    (tmp / "slim.json").write_text(json.dumps(slim))
    r = run(SCRIPTS / "fill_template.py", "minimal", tmp / "slim.json", tmp / "slim.html")
    html = (tmp / "slim.html").read_text()
    check(r.returncode == 0, "fill with optional fields missing")
    check("Selected Projects" not in html and "Languages" not in html, "empty sections removed")
    check("+46" not in html and "janedoe.dev" not in html, "empty contact items removed")
    check(not re.search(r'class="sep">·</span>\s*</div>', html), "no trailing separator")

    print("\n[no PDF engine: HTML + text fallback]")
    # Simulated: patch every renderer to report "unavailable"; nothing is uninstalled.
    fb_html = tmp / "fb_in.html"
    r = run(SCRIPTS / "fill_template.py", "minimal", TESTS / "sample_cv.json", fb_html)
    sim = (
        "import sys; sys.path.insert(0, %r); import build_cv as b; "
        "b.try_chrome = lambda *a, **k: False; b.try_weasyprint = lambda *a, **k: False; "
        "b.try_playwright = lambda *a, **k: False; "
        "out = sys.argv[2]; "
        "plain = b.build_cv(sys.argv[1], out); "
        "code = b.build_or_fallback(sys.argv[1], out); "
        "print('RESULT', plain, code)"
    ) % str(SCRIPTS)
    out_pdf = tmp / "fallback" / "Jane_Doe_CV.pdf"
    r = subprocess.run([sys.executable, "-c", sim, str(fb_html), str(out_pdf)], capture_output=True, text=True, cwd=tmp)
    check("RESULT False 2" in r.stdout, "no engine: plain build fails, --fallback path returns exit code 2")
    check(not out_pdf.exists(), "no PDF is produced")
    check(out_pdf.with_suffix(".html").exists() and "Jane Doe" in out_pdf.with_suffix(".html").read_text(), "finished HTML is delivered")
    check(out_pdf.with_suffix(".txt").exists() and "Jane Doe" in out_pdf.with_suffix(".txt").read_text(), "plain-text copy is delivered")
    check("Print" in r.stderr, "user is told how to save as PDF")

    print("\n[works from any working directory]")
    elsewhere = tmp / "elsewhere"
    elsewhere.mkdir()
    r = subprocess.run([sys.executable, str(SCRIPTS / "fill_template.py"), "minimal", str(TESTS / "sample_cv.json"), "w.html"], capture_output=True, text=True, cwd=elsewhere)
    check(r.returncode == 0 and (elsewhere / "w.html").exists(), "fill_template by name from another cwd")
    r = subprocess.run([sys.executable, str(SCRIPTS / "build_cv.py"), "w.html", "w.pdf", "--fallback"], capture_output=True, text=True, cwd=elsewhere)
    check(r.returncode in (0, 2) and ((elsewhere / "w.pdf").exists() or (elsewhere / "w.html").exists()), "build_cv --fallback from another cwd (exit %d)" % r.returncode)

    print("\n[skill frontmatter]")
    skill = (ROOT / "skills" / "made-for-you-cv" / "SKILL.md").read_text()
    m = re.match(r'---\nname: made-for-you-cv\ndescription: "(.*)"\n---\n', skill)
    check(bool(m), "frontmatter has name and quoted description")
    if m:
        d = m.group(1)
        check(len(d) < 900, f"description under 900 chars ({len(d)})")
        check("<" not in d and ">" not in d, "description has no angle brackets")
        for phrase in ("resume", "cover letter", "job application", "apply for this job", "tailor my CV", "personligt brev", "anpassa mitt CV"):
            check(phrase.lower() in d.lower(), f"description mentions '{phrase}'")
    check("python3 scripts/" not in skill, "SKILL.md has no cwd-relative script calls")

    print("\n[language: section headings, dates]")
    sv = dict(cv, LANG="sv")
    (tmp / "sv.json").write_text(json.dumps(sv, ensure_ascii=False))
    for tpl in ("minimal", "two-column", "designer-accent"):
        r = run(SCRIPTS / "fill_template.py", tpl, tmp / "sv.json", tmp / f"sv_{tpl}.html")
        h = (tmp / f"sv_{tpl}.html").read_text()
        heads = re.findall(r"<h2>(.*?)</h2>", h)
        check(r.returncode == 0 and "Erfarenhet" in heads and "Utbildning" in heads and "Språk" in heads and "Kompetenser" in heads,
              f"{tpl}: Swedish headings {heads}")
        check('lang="sv"' in h, f"{tpl}: html lang is sv")
        check("Experience" not in heads and "Education" not in heads, f"{tpl}: no English headings left")
    check("Profil" in (tmp / "sv_two-column.html").read_text() and "Kontakt" in (tmp / "sv_two-column.html").read_text(), "two-column: Profil and Kontakt")
    r = run(SCRIPTS / "fill_template.py", "minimal", TESTS / "sample_cv.json", tmp / "en.html")
    check("<h2>Experience</h2>" in (tmp / "en.html").read_text(), "default stays English")
    for lang, expected in (("sv", "5 oktober 2026"), ("de", "5. Oktober 2026"), ("en", "5 October 2026"), ("fr", "5 octobre 2026")):
        cl = json.loads((TESTS / "sample_cover_letter.json").read_text())
        cl.update(LANG=lang, DATE="2026-10-05")
        (tmp / "cl.json").write_text(json.dumps(cl))
        run(SCRIPTS / "fill_template.py", "cover-letter", tmp / "cl.json", tmp / "cl.html")
        check(expected in (tmp / "cl.html").read_text(), f"cover letter date localized for {lang} ({expected})")
    r = run(SCRIPTS / "fill_template.py", "minimal", tmp / "sv.json", tmp / "x.html")
    for lang in ("da", "nb", "de", "nl", "fr"):
        (tmp / "l.json").write_text(json.dumps(dict(cv, LANG=lang), ensure_ascii=False))
        r = run(SCRIPTS / "fill_template.py", "minimal", tmp / "l.json", tmp / "l.html")
        check(r.returncode == 0 and "<h2>Experience</h2>" not in (tmp / "l.html").read_text(), f"LANG={lang} localizes headings")

    print("\n[SECTION_ORDER]")
    for tpl in ("minimal", "two-column", "designer-accent"):
        (tmp / "o.json").write_text(json.dumps(dict(cv, SECTION_ORDER="education,projects,experience")))
        run(SCRIPTS / "fill_template.py", tpl, tmp / "o.json", tmp / "o.html")
        heads = re.findall(r"<h2>(.*?)</h2>", (tmp / "o.html").read_text())
        e, pj, x = heads.index("Education"), next(i for i, h in enumerate(heads) if "Projects" in h), heads.index("Experience")
        check(e < pj < x, f"{tpl}: education, projects, then experience {heads}")
    (tmp / "o.json").write_text(json.dumps(dict(cv, SECTION_ORDER=["education"])))
    run(SCRIPTS / "fill_template.py", "minimal", tmp / "o.json", tmp / "o.html")
    heads = re.findall(r"<h2>(.*?)</h2>", (tmp / "o.html").read_text())
    check(heads[:2] == ["Education", "Experience"], f"partial order keeps the rest in default order {heads}")

    print("\n[Word .docx]")
    try:
        import docx  # noqa: F401
        have_docx = True
    except ImportError:
        have_docx = False
        print("  SKIP python-docx not installed (CI installs it)")
    if have_docx:
        from docx import Document
        (tmp / "d.json").write_text(json.dumps(dict(cv, LANG="sv", SECTION_ORDER="education,projects,experience"), ensure_ascii=False))
        out_docx = tmp / "cv.docx"
        r = run(SCRIPTS / "build_docx.py", tmp / "d.json", out_docx)
        check(r.returncode == 0 and out_docx.exists(), "build_docx writes a .docx " + r.stderr.strip()[-120:])
        d = Document(str(out_docx))
        heads = [p.text for p in d.paragraphs if p.style.name == "Heading 1"]
        check(heads == ["Profil", "Utbildning", "Projekt", "Erfarenhet", "Kompetenser", "Språk"], f"headings in order {heads}")
        check(len(d.tables) == 0 and len(d.inline_shapes) == 0, "no tables or images")
        check(all(not s.header.paragraphs or not any(p.text for p in s.header.paragraphs) for s in d.sections)
              and all(not any(p.text for p in s.footer.paragraphs) for s in d.sections), "no header or footer content")
        text = "\n".join(p.text for p in d.paragraphs)
        check("Jane Doe" in text and "Migrated a 40-page legacy app" in text and "Frontend: React" in text, "name, bullets and skills are present")
        check(sum(1 for p in d.paragraphs if p.style.name == "List Bullet") >= 4, "bullets are real list paragraphs")
        cl = json.loads((TESTS / "sample_cover_letter.json").read_text())
        (tmp / "cld.json").write_text(json.dumps(cl))
        r = run(SCRIPTS / "build_docx.py", tmp / "cld.json", tmp / "cl.docx")
        check(r.returncode == 0 and "Dear Anna Lindqvist" in "\n".join(p.text for p in Document(str(tmp / "cl.docx")).paragraphs), "cover letter .docx")
    # python-docx missing and not installable: skip with a note, exit 0, no file.
    stub = (
        "import sys; sys.path.insert(0, %r); import build_docx as b; "
        "b.load_docx = lambda: None; import json; "
        "r = b.build(json.load(open(sys.argv[1])), sys.argv[2]); print('RESULT', r)"
    ) % str(SCRIPTS)
    r = subprocess.run([sys.executable, "-c", stub, str(TESTS / "sample_cv.json"), str(tmp / "none.docx")], capture_output=True, text=True)
    check("RESULT None" in r.stdout and "Skipped .docx" in r.stderr and not (tmp / "none.docx").exists(), "missing python-docx: one-line skip, no file")

    print("\n[--fit]")
    over = dict(cv, EXPERIENCE_BLOCK=cv["EXPERIENCE_BLOCK"] * 2)
    (tmp / "over.json").write_text(json.dumps(over))
    for tpl in ("minimal", "designer-accent"):
        run(SCRIPTS / "fill_template.py", tpl, tmp / "over.json", tmp / f"over_{tpl}.html")
        r = run(SCRIPTS / "build_cv.py", tmp / f"over_{tpl}.html", tmp / f"over_{tpl}_plain.pdf")
        plain = re.search(r"Page count: (\d+)", r.stdout)
        r = run(SCRIPTS / "build_cv.py", tmp / f"over_{tpl}.html", tmp / f"over_{tpl}_fit.pdf", "--fit")
        fitted = re.search(r"Page count: (\d+)", r.stdout)
        check(bool(plain) and int(plain.group(1)) == 2, f"{tpl}: sample is 2 pages without --fit")
        check(bool(fitted) and int(fitted.group(1)) == 1 and "Fit to 1 page" in r.stdout, f"{tpl}: --fit brings it to 1 page and reports what it did")
    sys.path.insert(0, str(SCRIPTS))
    import build_cv as bc
    ok = True
    for tpl in ("minimal", "two-column", "designer-accent"):
        src = (ROOT / "skills" / "made-for-you-cv" / "assets" / "templates" / f"{tpl}.html").read_text()
        before = [float(x) for x in re.findall(r"font-size:\s*([\d.]+)pt", src)]
        for step in bc.FIT_STEPS:
            after = [float(x) for x in re.findall(r"font-size:\s*([\d.]+)pt", bc.tighten_html(src, step))]
            ok &= all(a >= min(b, 9.5) for a, b in zip(after, before))
    check(ok, "--fit never takes text at or above 9.5pt below 9.5pt")
    r = run(SCRIPTS / "fill_template.py", "minimal", TESTS / "sample_cv.json", tmp / "fine.html")
    r = run(SCRIPTS / "build_cv.py", tmp / "fine.html", tmp / "fine.pdf", "--fit")
    check("Fit:" not in r.stdout and "Page count: 1" in r.stdout, "--fit leaves a 1-page CV alone")

    print("\n[eval fixtures are fictional]")
    evals = TESTS / "evals"
    cases = sorted(evals.glob("case_*"))
    check(len(cases) == 3, f"3 eval cases ({len(cases)})")
    for c in cases:
        meta = json.loads((c / "case.json").read_text())
        text = (c / "cv.md").read_text() + (c / "job_ad.md").read_text()
        emails = set(re.findall(r"[\w.+-]+@[\w.-]+\.\w+", text))
        phones = re.findall(r"\b0\d{2,4}[ -]?\d{3}[ -]?\d{2,3}[ -]?\d{0,3}\b", text)
        check(all(e.endswith("@example.com") for e in emails) and emails, f"{c.name}: emails use example.com")
        check(all(re.search(r"0{3}|00 ?\d{3}", p) for p in phones), f"{c.name}: phone numbers are 07x-000 style")
        check(all(k in meta for k in ("lang", "pages", "forbidden_personal", "forbidden_skills", "must_keywords")), f"{c.name}: case.json is complete")
        check((c / "prompt.txt").exists(), f"{c.name}: prompt.txt")

    print("\n[guards]")
    (tmp / "bad.json").write_text(json.dumps({"NAME": "X"}))
    r = run(SCRIPTS / "fill_template.py", "minimal", tmp / "bad.json", tmp / "bad.html")
    check(r.returncode != 0, "missing required fields rejected")
    (tmp / "raw.html").write_text("<html>{{NAME}}</html>")
    r = run(SCRIPTS / "build_cv.py", tmp / "raw.html", tmp / "raw.pdf")
    check(r.returncode != 0, "build refuses unfilled placeholders")

print()
if failures:
    print(f"{len(failures)} check(s) failed")
    sys.exit(1)
print("All checks passed")
