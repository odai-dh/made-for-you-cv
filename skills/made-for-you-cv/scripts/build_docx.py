#!/usr/bin/env python3
"""
build_docx.py: build a plain, ATS-friendly Word (.docx) CV or cover letter from
the same JSON that fill_template.py uses.

Usage:
    build_docx.py <data.json> <output.docx>

Designed for recruitment portals and applicant tracking systems: one column,
standard Heading 1 section headings, real bullet lists, no tables, text boxes,
images, or headers/footers. Section headings follow LANG and SECTION_ORDER just
like the HTML templates.

Needs python-docx. If it is missing the script tries `pip install python-docx`
once; if that fails it prints a one-line note and exits 0 without a file, so the
rest of the delivery can go ahead. Exit codes: 0 built or skipped (see output),
1 error in the input.
"""

import json
import re
import subprocess
import sys
from html.parser import HTMLParser
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import cv_common

FONT = "Calibri"


def load_docx():
    try:
        import docx
        return docx
    except ImportError:
        pass
    try:
        subprocess.run(
            [sys.executable, "-m", "pip", "install", "--quiet", "python-docx"],
            capture_output=True, timeout=180, check=True,
        )
        import docx
        return docx
    except Exception:
        return None


class Blocks(HTMLParser):
    """Parse the HTML blocks the templates use into simple dicts."""

    VOID = {"br", "img", "hr"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.entries, self.groups, self.spans, self.paras = [], [], [], []
        self._stack = []  # (tag, class names, text buffer)
        self._entry = None
        self._group = None

    def handle_starttag(self, tag, attrs):
        if tag in self.VOID:
            if tag == "br" and self._stack:
                self._stack[-1][2].append(" ")
            return
        names = (dict(attrs).get("class", "") or "").split()
        if tag == "div" and names[:1] == ["entry"]:
            self._entry = {"title": "", "meta": "", "sub": "", "bullets": [], "tech": ""}
            self.entries.append(self._entry)
        elif tag == "div" and "skills-group" in names:
            self._group = {"label": "", "items": ""}
            self.groups.append(self._group)
        self._stack.append((tag, names, []))

    def handle_data(self, data):
        if self._stack:
            self._stack[-1][2].append(data)

    def handle_endtag(self, tag):
        if tag in self.VOID or not self._stack:
            return
        # Close back to the matching open tag (tolerates sloppy HTML).
        while self._stack and self._stack[-1][0] != tag:
            self._stack.pop()
        if not self._stack:
            return
        _, names, buf = self._stack.pop()
        text = re.sub(r"\s+", " ", "".join(buf)).strip()
        if self._stack:
            self._stack[-1][2].append(" " + text + " ")  # bubble text up to the parent
        if self._entry is not None:
            if "entry-title" in names:
                self._entry["title"] = text
            elif "entry-meta" in names:
                self._entry["meta"] = text
            elif "entry-sub" in names:
                self._entry["sub"] = text
            elif "tech" in names:
                self._entry["tech"] = text
            elif tag == "li" and text:
                self._entry["bullets"].append(text)
        if self._group is not None:
            if tag == "strong":
                self._group["label"] = text
            elif "items" in names:
                self._group["items"] = text
        if tag == "span" and self._entry is None and self._group is None and text and (not names or "lang" in names):
            self.spans.append(text)
        if tag == "p" and text:
            self.paras.append(text)
        if tag == "div" and names[:1] == ["entry"]:
            self._entry = None
        if tag == "div" and "skills-group" in names:
            self._group = None


def parse(block):
    p = Blocks()
    p.feed(block or "")
    return p


def build(data, out_path):
    docx = load_docx()
    if docx is None:
        print("Skipped .docx: python-docx is not available and could not be installed here.", file=sys.stderr)
        return None
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: F401

    code, _ = cv_common.normalize_lang(data.get("LANG"))
    labels = cv_common.labels_for(code)
    labels.update({k: v for k, v in data.items() if k.startswith("H_") and v})

    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    h1 = doc.styles["Heading 1"]
    h1.font.name = FONT
    h1.font.size = Pt(12)
    h1.font.bold = True
    h1.font.color.rgb = docx.shared.RGBColor(0x1A, 0x1A, 0x1A)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(3)

    def para(text="", bold=False, italic=False, size=None, style=None, after=None):
        p = doc.add_paragraph(style=style)
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    # Header: name, title, contact line (plain paragraphs, no header/footer area).
    para(data.get("NAME", ""), bold=True, size=18, after=0)
    if data.get("TITLE"):
        para(data["TITLE"], size=11, after=2)
    contact = [data.get(k, "") for k in ("LOCATION", "EMAIL", "PHONE")]
    for k in ("PORTFOLIO", "GITHUB", "LINKEDIN"):
        v = data.get(k + "_DISPLAY") or data.get(k, "")
        contact.append(re.sub(r"^[a-z]+://", "", v) if v else "")
    line = " | ".join(c for c in contact if c)
    if line:
        para(line, size=9.5, after=6)

    if "BODY" in data:  # cover letter
        if data.get("DATE"):
            para(cv_common.localize_date(data["DATE"], code), after=8)
        for l in (data.get("RECIPIENT_BLOCK") or "").splitlines():
            if l.strip():
                para(re.sub(r"<[^>]+>", "", l), after=0)
        para(data.get("SALUTATION", ""), after=6)
        for t in parse(data["BODY"]).paras or [re.sub(r"<[^>]+>", "", data["BODY"])]:
            para(t, after=6)
        para(data.get("CLOSING", ""), after=14)
        para(data.get("NAME", ""), bold=True)
        doc.save(str(out_path))
        return out_path

    if data.get("SUMMARY"):
        doc.add_paragraph(labels["H_PROFILE"], style="Heading 1")
        summ = data["SUMMARY"]
        para(parse(summ).paras[0] if "<p" in summ and parse(summ).paras else re.sub(r"<[^>]+>", "", summ))

    def entries_section(name, block):
        parsed = parse(block)
        if not parsed.entries:
            return
        doc.add_paragraph(labels[name], style="Heading 1")
        for e in parsed.entries:
            line1 = e["title"]
            para(line1, bold=True, after=0)
            sub = " | ".join(x for x in (e["sub"], e["meta"]) if x)
            if sub:
                para(sub, italic=True, size=10, after=2)
            for b in e["bullets"]:
                para(b, style="List Bullet", after=1)
            if e["tech"]:
                para(e["tech"], size=9.5, after=6)
            else:
                doc.paragraphs[-1].paragraph_format.space_after = Pt(6)

    def skills_section(block):
        parsed = parse(block)
        if not parsed.groups:
            text = re.sub(r"<[^>]+>", " ", block or "")
            text = re.sub(r"\s+", " ", text).strip()
            if not text:
                return
            doc.add_paragraph(labels["H_SKILLS"], style="Heading 1")
            para(text)
            return
        doc.add_paragraph(labels["H_SKILLS"], style="Heading 1")
        for g in parsed.groups:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            if g["label"]:
                p.add_run(g["label"] + ": ").bold = True
            p.add_run(g["items"])

    def languages_section(block):
        items = parse(block).spans
        if not items:
            text = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", block or "")).strip()
            items = [text] if text else []
        if not items:
            return
        doc.add_paragraph(labels["H_LANGUAGES"], style="Heading 1")
        para(", ".join(items))

    for name in cv_common.full_order(data.get("SECTION_ORDER", "")):
        if name == "experience":
            entries_section("H_EXPERIENCE", data.get("EXPERIENCE_BLOCK", ""))
        elif name == "projects":
            entries_section("H_PROJECTS", data.get("PROJECTS_BLOCK", ""))
        elif name == "education":
            entries_section("H_EDUCATION", data.get("EDUCATION_BLOCK", ""))
        elif name == "skills":
            skills_section(data.get("SKILLS_BLOCK", ""))
        elif name == "languages":
            languages_section(data.get("LANGUAGES_BLOCK", ""))

    doc.core_properties.title = f"{data.get('NAME', '')} - CV".strip(" -")
    doc.core_properties.author = data.get("NAME", "")
    doc.save(str(out_path))
    return out_path


def main():
    if len(sys.argv) != 3:
        print("Usage: build_docx.py <data.json> <output.docx>", file=sys.stderr)
        sys.exit(1)
    try:
        data = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    except Exception as e:
        print(f"❌ Could not read JSON {sys.argv[1]}: {e}", file=sys.stderr)
        sys.exit(1)
    for k in ("NAME", "EMAIL"):
        if not data.get(k):
            print(f"❌ Missing required field: {k}", file=sys.stderr)
            sys.exit(1)
    out = Path(sys.argv[2])
    out.parent.mkdir(parents=True, exist_ok=True)
    result = build(data, out)
    if result:
        print(f"✅ Word document written: {out}")


if __name__ == "__main__":
    main()
